"""
Excel storage with professional formatting:
- Frozen header row
- Auto-filters
- Readable column widths
- Text wrapping for descriptions
- Multiple sheets (Dataset, Summary, Data Dictionary, Verification Log)
"""
import pandas as pd
from typing import List, Dict, Any
from pathlib import Path
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import logging

logger = logging.getLogger(__name__)


class ExcelStorage:
    @staticmethod
    def save(filename: str, records: list, summary_data: dict,
             verification_log: list):
        output_dir = Path("data/output")
        output_dir.mkdir(parents=True, exist_ok=True)
        filepath = output_dir / filename

        # ── Flatten records ──
        flat_records = []
        for r in records:
            d = r.dict() if hasattr(r, 'dict') else r.model_dump()
            content = d.pop('content', {})
            source = d.pop('source', {})
            verification = d.pop('verification', {})
            quality = d.pop('quality', {})
            flat = {
                'id': d.get('id'),
                'name': content.get('name'),
                'company_or_developer': content.get('company_or_developer'),
                'description': content.get('description'),
                'primary_task': content.get('primary_task'),
                'categories': ', '.join(content.get('categories', [])),
                'pricing_model': content.get('pricing_model'),
                'is_open_source': content.get('is_open_source'),
                'official_website': content.get('official_website'),
                'inputs': content.get('inputs'),
                'outputs': content.get('outputs'),
                'supported_platforms': content.get('supported_platforms'),
                'has_api': content.get('has_api'),
                'logo_url': content.get('logo_url'),
                'status': content.get('status'),
                'launch_date': content.get('launch_date'),
                'last_verified': content.get('last_verified'),
                'source_name': source.get('name'),
                'source_url': source.get('url'),
                'collected_at': d.get('collectedAt'),
            }
            flat_records.append(flat)

        df = pd.DataFrame(flat_records)

        # ── Data Dictionary ──
        dict_entries = []
        if flat_records:
            for key in flat_records[0].keys():
                sample_vals = [r.get(key) for r in flat_records[:5] if r.get(key)]
                sample = str(sample_vals[0])[:80] if sample_vals else ""
                dict_entries.append({
                    "Field": key,
                    "Type": type(sample_vals[0]).__name__ if sample_vals else "str",
                    "Description": _field_description(key),
                    "Sample": sample,
                })
        df_dict = pd.DataFrame(dict_entries)

        # ── Summary ──
        summary_rows = [{"Metric": k, "Value": v} for k, v in summary_data.items()]
        df_summary = pd.DataFrame(summary_rows)

        # ── Verification Log ──
        df_log = pd.DataFrame(verification_log) if verification_log else pd.DataFrame(
            [{"status": "No verification logs generated"}])

        # ── Write workbook ──
        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='AI Tools Dataset', index=False)
            df_summary.to_excel(writer, sheet_name='Dataset Summary', index=False)
            df_dict.to_excel(writer, sheet_name='Data Dictionary', index=False)
            df_log.to_excel(writer, sheet_name='Verification Log', index=False)

            wb = writer.book

            # ── Format the main dataset sheet ──
            ws = wb['AI Tools Dataset']
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions

            # Header style
            header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True, size=11)
            thin_border = Border(
                bottom=Side(style='thin', color='CCCCCC')
            )

            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')

            # Column widths & text wrapping
            col_widths = {
                'id': 12, 'name': 25, 'company_or_developer': 22,
                'description': 55, 'primary_task': 20, 'categories': 30,
                'pricing_model': 14, 'is_open_source': 13,
                'official_website': 35, 'inputs': 35, 'outputs': 35,
                'supported_platforms': 18, 'has_api': 10, 'logo_url': 30,
                'status': 10, 'launch_date': 14, 'last_verified': 14,
                'source_name': 18, 'source_url': 35, 'collected_at': 22,
            }
            wrap_cols = {'description', 'inputs', 'outputs', 'categories'}

            for col_idx, col_name in enumerate(df.columns, 1):
                letter = get_column_letter(col_idx)
                ws.column_dimensions[letter].width = col_widths.get(col_name, 15)

                if col_name in wrap_cols:
                    for row in ws.iter_rows(min_row=2, min_col=col_idx,
                                            max_col=col_idx, max_row=ws.max_row):
                        for cell in row:
                            cell.alignment = Alignment(wrap_text=True, vertical='top')

            # Data rows border
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row,
                                    max_col=ws.max_column):
                for cell in row:
                    cell.border = thin_border

            # ── Format summary sheet ──
            ws_sum = wb['Dataset Summary']
            ws_sum.column_dimensions['A'].width = 30
            ws_sum.column_dimensions['B'].width = 50
            for cell in ws_sum[1]:
                cell.font = Font(bold=True)

        logger.info(f"Excel workbook saved: {filepath}")


def _field_description(field: str) -> str:
    """Return human-readable field descriptions for the data dictionary."""
    descriptions = {
        'id': 'Deterministic unique identifier for the tool',
        'name': 'Official name of the AI tool',
        'company_or_developer': 'Company or developer that created the tool',
        'description': 'Concise factual description of what the tool does',
        'primary_task': 'Primary task category from controlled taxonomy',
        'categories': 'Comma-separated list of categories',
        'pricing_model': 'Pricing model: FREE, FREEMIUM, PAID, ENTERPRISE',
        'is_open_source': 'Whether the tool is open source',
        'official_website': 'Verified official website URL',
        'inputs': 'What the tool accepts as input',
        'outputs': 'What the tool produces as output',
        'supported_platforms': 'Supported platforms (web, mobile, desktop, API)',
        'has_api': 'Whether the tool provides an API',
        'logo_url': 'URL to official logo (verified, not fabricated)',
        'status': 'Current status: Active, Beta, Deprecated',
        'launch_date': 'Date the tool was launched',
        'last_verified': 'Date the record was last verified',
        'source_name': 'Discovery source name',
        'source_url': 'Discovery source URL',
        'collected_at': 'ISO-8601 timestamp when data was collected',
    }
    return descriptions.get(field, field.replace('_', ' ').title())
